from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
DOCS_DIR = ROOT / "docs"
ASSET_DIR = DOCS_DIR / "assets"
OUTPUT_PATH = DOCS_DIR / "ReturnWise_MAS_Technical_Report.docx"

BLUE = "1F4E79"
LIGHT_BLUE = "D9EAF7"
PALE_GRAY = "F3F6F8"
TEXT = "1F2933"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text: str, bold: bool = False, size: int = 9) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(TEXT)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def add_table(document: Document, headers: list[str], rows: list[list[str]], widths: list[float]) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        set_cell_shading(cell, LIGHT_BLUE)
        set_cell_text(cell, header, bold=True, size=9)
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            set_cell_text(cells[index], value, size=8)
    for row in table.rows:
        for index, width in enumerate(widths):
            row.cells[index].width = Inches(width)
            for paragraph in row.cells[index].paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
    document.add_paragraph()


def add_footer(section) -> None:
    footer = section.footer
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("ReturnWise MAS Technical Report")
    run.font.name = "Times New Roman"
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string("666666")


def configure_styles(document: Document) -> None:
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(TEXT)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(6)

    for style_name, size in [("Heading 1", 15), ("Heading 2", 12), ("Heading 3", 11)]:
        style = styles[style_name]
        style.font.name = "Times New Roman"
        style.font.bold = True
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(BLUE)
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(4)


def add_heading(document: Document, text: str, level: int = 1) -> None:
    document.add_heading(text, level=level)


def add_body(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(text)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY


def add_bullets(document: Document, items: list[str]) -> None:
    for item in items:
        paragraph = document.add_paragraph(style="List Bullet")
        run = paragraph.add_run(item)
        run.font.name = "Times New Roman"
        run.font.size = Pt(10.5)


def add_numbered(document: Document, items: list[str]) -> None:
    for item in items:
        paragraph = document.add_paragraph(style="List Number")
        run = paragraph.add_run(item)
        run.font.name = "Times New Roman"
        run.font.size = Pt(10.5)


def add_code_box(document: Document, lines: list[str]) -> None:
    table = document.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.rows[0].cells[0]
    set_cell_shading(cell, PALE_GRAY)
    cell.text = ""
    for index, line in enumerate(lines):
        paragraph = cell.paragraphs[0] if index == 0 else cell.add_paragraph()
        run = paragraph.add_run(line)
        run.font.name = "Courier New"
        run.font.size = Pt(8.5)
        paragraph.paragraph_format.space_after = Pt(0)
    document.add_paragraph()


def create_architecture_diagram(path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    image = Image.new("RGB", (1500, 760), "white")
    draw = ImageDraw.Draw(image)
    try:
        title_font = ImageFont.truetype("arialbd.ttf", 38)
        font = ImageFont.truetype("arial.ttf", 28)
        small_font = ImageFont.truetype("arial.ttf", 22)
    except OSError:
        title_font = ImageFont.load_default()
        font = ImageFont.load_default()
        small_font = ImageFont.load_default()

    draw.text((45, 35), "ReturnWise MAS Workflow", fill="#1F4E79", font=title_font)

    nodes = [
        ("Local JSON\nRequests", 65, 210, 260, 120, "#E9F5EC"),
        ("Intake\nAgent", 345, 210, 220, 120, "#D9EAF7"),
        ("Policy\nAgent", 640, 105, 220, 120, "#D9EAF7"),
        ("Risk\nAgent", 640, 315, 220, 120, "#D9EAF7"),
        ("Resolution\nAgent", 930, 210, 245, 120, "#D9EAF7"),
        ("Reports\nJSON + MD", 1255, 210, 210, 120, "#F7E7D2"),
        ("Policy\nMarkdown", 345, 505, 220, 105, "#E9F5EC"),
        ("WorkflowState\nShared Context", 930, 505, 245, 105, "#F3F6F8"),
        ("JSONL\nTrace Log", 1255, 505, 210, 105, "#F3F6F8"),
    ]

    for label, x, y, w, h, fill in nodes:
        draw.rounded_rectangle((x, y, x + w, y + h), radius=20, fill=fill, outline="#7A8793", width=3)
        lines = label.split("\n")
        total_height = len(lines) * 32
        for i, line in enumerate(lines):
            bbox = draw.textbbox((0, 0), line, font=font if i == 0 else small_font)
            draw.text((x + (w - (bbox[2] - bbox[0])) / 2, y + (h - total_height) / 2 + i * 35), line, fill="#1F2933", font=font if i == 0 else small_font)

    arrows = [
        ((260, 270), (345, 270)),
        ((565, 270), (640, 165)),
        ((565, 270), (640, 375)),
        ((860, 165), (930, 270)),
        ((860, 375), (930, 270)),
        ((1175, 270), (1255, 270)),
        ((455, 505), (700, 225)),
        ((1052, 330), (1052, 505)),
        ((1175, 558), (1255, 558)),
    ]
    for start, end in arrows:
        draw.line((start, end), fill="#1F4E79", width=4)
        draw.ellipse((end[0] - 7, end[1] - 7, end[0] + 7, end[1] + 7), fill="#1F4E79")

    draw.text((45, 680), "All data, tools, logs, and model calls run locally. Ollama provides local SLM support.", fill="#4B5563", font=small_font)
    image.save(path)


def add_title_page(document: Document) -> None:
    for _ in range(3):
        document.add_paragraph()
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("ReturnWise MAS")
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(24)
    run.font.color.rgb = RGBColor.from_string(BLUE)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("A Locally Hosted Multi-Agent System for E-Commerce Return Triage")
    run.font.name = "Times New Roman"
    run.font.size = Pt(14)
    run.italic = True

    document.add_paragraph()
    meta_rows = [
        ["Module", "SE4010 - Current Trends in Software Engineering"],
        ["Assessment", "Assignment 2 - Machine Learning"],
        ["Institution", "Sri Lanka Institute of Information Technology"],
        ["Technology Stack", "Python, LangGraph, Ollama, local JSON/Markdown tools"],
        ["Repository", "Add GitHub repository link here after pushing"],
    ]
    add_table(document, ["Item", "Details"], meta_rows, [1.45, 4.85])

    add_heading(document, "Abstract", 1)
    add_body(
        document,
        "This report presents ReturnWise MAS, a zero-cost local multi-agent system built for e-commerce return request triage. "
        "The system uses four coordinated agents to validate requests, apply return-policy rules, assess operational risk, and generate final support decisions. "
        "The implementation demonstrates local small language model usage through Ollama while keeping business-critical calculations inside auditable Python tools. "
        "The design also includes typed global state, JSONL observability logs, automated tests, and a deterministic evaluation harness.",
    )
    document.add_page_break()


def add_static_contents(document: Document) -> None:
    add_heading(document, "Table of Contents", 1)
    contents = [
        "1. Introduction and Problem Domain",
        "2. System Architecture and Workflow",
        "3. Agent Design",
        "4. Custom Tool Integration",
        "5. State Management and Observability",
        "6. Evaluation Methodology",
        "7. Individual Contribution Proof",
        "8. Conclusion",
    ]
    for item in contents:
        paragraph = document.add_paragraph(item)
        paragraph.paragraph_format.left_indent = Inches(0.2)
        paragraph.paragraph_format.space_after = Pt(2)
    document.add_page_break()


def build_document() -> None:
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    diagram_path = ASSET_DIR / "returnwise_architecture.png"
    create_architecture_diagram(diagram_path)

    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    configure_styles(document)
    add_footer(section)

    add_title_page(document)
    add_static_contents(document)

    add_heading(document, "1. Introduction and Problem Domain", 1)
    add_body(
        document,
        "Online stores handle return requests that differ widely in complexity. Some requests only require a basic policy check, while others involve defective products, damaged deliveries, high-value used items, or repeated return behaviour. "
        "ReturnWise MAS addresses this problem by automating the first review stage. It does not replace a human support team; it prepares consistent decisions, supporting notes, and traceable evidence for review.",
    )
    add_body(
        document,
        "The selected domain is suitable for an agentic system because the task naturally separates into intake, policy interpretation, risk analysis, and final resolution. "
        "Each step depends on earlier context, but each also has its own responsibility and tool requirements. The project therefore demonstrates a multi-agent workflow rather than a single conversational chatbot.",
    )
    add_bullets(
        document,
        [
            "Input data is stored locally as JSON return requests.",
            "Policy knowledge is stored locally as a markdown policy document.",
            "The language model runs locally through Ollama, with no paid API keys.",
            "Outputs are saved as markdown, JSON, and JSONL trace logs.",
        ],
    )

    add_heading(document, "2. System Architecture and Workflow", 1)
    add_body(
        document,
        "ReturnWise uses a controlled pipeline orchestrated through LangGraph. The graph contains four agent nodes: Intake, Policy, Risk, and Resolution. "
        "When LangGraph is available, the implementation compiles the workflow as a StateGraph. A small sequential fallback is included only to keep local tests runnable before dependencies are installed.",
    )
    figure = document.add_paragraph()
    figure.alignment = WD_ALIGN_PARAGRAPH.CENTER
    figure.add_run().add_picture(str(diagram_path), width=Inches(6.45))
    caption = document.add_paragraph("Figure 1: ReturnWise MAS local multi-agent architecture.")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.runs[0].italic = True
    caption.runs[0].font.size = Pt(9)
    add_numbered(
        document,
        [
            "The Intake Agent reads and normalizes return requests.",
            "The Policy Agent compares each request with the local return policy.",
            "The Risk Agent scores operational risk using deterministic rules.",
            "The Resolution Agent combines policy and risk outputs and writes final decisions.",
        ],
    )

    add_heading(document, "3. Agent Design", 1)
    add_body(
        document,
        "The agents are designed with narrow prompts and clearly separated responsibilities. This is important because local small language models can lose focus if asked to perform broad, loosely defined tasks. "
        "Business rules are handled by tools, while the model is used for short audit notes and wording support.",
    )
    add_table(
        document,
        ["Agent", "Prompt Focus", "Responsibility", "Primary Tool"],
        [
            ["Intake Agent", "Use only provided fields; do not invent missing facts.", "Validate and normalize raw request data.", "load_return_requests; normalize_return_request"],
            ["Policy Agent", "Use the local policy text only.", "Check eligibility, policy window, and suggested status.", "read_policy_text; match_policy"],
            ["Risk Agent", "Describe risk as internal operational signals.", "Score risk from value, condition, return history, and policy result.", "calculate_risk_score"],
            ["Resolution Agent", "Produce polite customer messages and clear internal notes.", "Generate approve, replace, escalate, or reject decisions.", "build_decision; write_decision_outputs"],
        ],
        [1.05, 1.6, 2.15, 1.55],
    )
    add_body(
        document,
        "The interaction strategy is sequential because return triage has a natural order: data must be valid before policy can be applied, and policy results must be available before risk and final resolution. "
        "This structure keeps handoffs easy to inspect and reduces overlap between agents.",
    )

    add_heading(document, "4. Custom Tool Integration", 1)
    add_body(
        document,
        "The project uses custom Python tools with type hints, docstrings, validation, and explicit errors. These tools give agents access to local files, policy rules, date calculations, risk scoring, and report writing. "
        "This satisfies the requirement that agents interact with the environment rather than depending only on model knowledge.",
    )
    add_table(
        document,
        ["Tool", "Purpose"],
        [
            ["load_return_requests", "Reads a local JSON list of return requests."],
            ["normalize_return_request", "Validates required fields, converts types, and checks dates."],
            ["read_policy_text", "Loads the local markdown return policy."],
            ["match_policy", "Applies policy windows and returns eligibility with a reason."],
            ["calculate_risk_score", "Creates a 0-100 risk score and internal risk flags."],
            ["build_decision", "Combines policy and risk into a final support decision."],
            ["write_decision_outputs", "Writes final decisions to markdown and JSON files."],
        ],
        [2.0, 4.25],
    )
    add_code_box(
        document,
        [
            'policy_text = read_policy_text("sample_data/return_policy.md")',
            "policy_match = match_policy(request, policy_text)",
            "risk = calculate_risk_score(request, policy_match)",
            "decision = build_decision(request, policy_match, risk)",
        ],
    )

    add_heading(document, "5. State Management and Observability", 1)
    add_body(
        document,
        "Global state is represented by the typed WorkflowState structure in the source code. Each agent receives the state, adds its own output, and passes the updated state to the next agent. "
        "The shared state prevents context loss because final decisions can be traced back to normalized inputs, policy matches, and risk assessments.",
    )
    add_table(
        document,
        ["State Field", "Purpose"],
        [
            ["requests", "Normalized return requests produced by the Intake Agent."],
            ["policy_matches", "Eligibility, window, and policy reason from the Policy Agent."],
            ["risk_assessments", "Risk level, score, and flags from the Risk Agent."],
            ["decisions", "Final support outcomes produced by the Resolution Agent."],
            ["agent_notes", "Short audit notes written during execution."],
            ["log_path", "Location of the JSONL trace file for the run."],
        ],
        [1.8, 4.45],
    )
    add_body(
        document,
        "Observability is implemented through an append-only JSONL logger. The log records agent start events, tool calls, tool results, agent outputs, and model errors. "
        "This gives a practical AgentOps view of each run and makes the workflow easier to debug during demonstrations.",
    )

    add_heading(document, "6. Evaluation Methodology", 1)
    add_body(
        document,
        "The evaluation strategy combines unit tests with an end-to-end harness. The harness runs with local deterministic settings by disabling optional Ollama calls, which makes the test reliable on any machine. "
        "The normal application mode still supports Ollama for local SLM-based audit notes.",
    )
    add_table(
        document,
        ["Evaluation Area", "Validation"],
        [
            ["Policy accuracy", "Defective electronics and policy windows are checked with assertions."],
            ["Risk reliability", "Risk scores must remain within the 0-100 range and identify high-risk cases."],
            ["Output quality", "The system must produce four decisions with valid statuses and messages."],
            ["Security constraint", "The harness checks that paid cloud API keys are not required."],
            ["Observability", "Tests verify that JSONL logs include agent and tool-call events."],
        ],
        [1.75, 4.5],
    )
    add_body(
        document,
        "The verified commands are `python tests/evaluation_harness.py` and `python -m unittest discover -s tests -p \"test_*.py\"`. Both are suitable for the project demo because they can be run locally without network access.",
    )

    add_heading(document, "7. Individual Contribution Proof", 1)
    add_body(document, "This submission was completed by a single contributor.")
    add_table(
        document,
        ["Student", "Agent", "Tool", "Test Contribution"],
        [
            ["Abishek S (IT22186942)", "Intake Agent, Policy Agent, Risk Agent, Resolution Agent", "load_return_requests; normalize_return_request; read_policy_text; match_policy; calculate_risk_score; build_decision; write_decision_outputs", "Missing-field, validation, policy, risk, and end-to-end output tests."],
        ],
        [1.45, 2.4, 4.2, 2.2],
    )
    add_body(
        document,
        "The main challenge was balancing local SLM use with reliability. The solution was to use Ollama for concise reasoning notes while keeping policy checks, risk scoring, and file writing inside deterministic tools.",
    )

    add_heading(document, "8. Conclusion", 1)
    add_body(
        document,
        "ReturnWise MAS demonstrates a complete locally hosted multi-agent system with four agents, custom tools, LangGraph orchestration, typed state management, JSONL observability, local Ollama support, and automated evaluation. "
        "Every decision can be traced through the workflow, making the result practical for review.",
    )

    document.save(OUTPUT_PATH)


if __name__ == "__main__":
    build_document()
    print(OUTPUT_PATH)
